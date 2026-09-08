"""
Import massal Purchase Order dari Excel.
Baris dengan `grup` sama -> 1 PO dengan banyak item, dibuat dalam 1 transaksi.
FK (entitas/suplier/produk) di-match berdasarkan field `nama`.

CATATAN: field selain yang dipakai eksplisit di sini (mis. field lain
di DiauditModel selain dibuat_oleh/dibuat_pada) diasumsikan punya
default/auto_now_add. Kalau ada field wajib lain tanpa default,
IntegrityError akan muncul dengan nama field aslinya di pesan galat.
"""
from decimal import Decimal, InvalidOperation
from datetime import datetime, date

from django.db import transaction, IntegrityError
from openpyxl import load_workbook

from core.models import Entitas
from master.models import Suplier, Produk
from .pembelian import PurchaseOrder, PurchaseOrderItem, StatusPO

KOLOM_WAJIB = {'grup', 'entitas', 'suplier', 'tanggal', 'produk', 'qty_pesan', 'harga_per_kg'}
FORMAT_TANGGAL = ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y')


class BarisGalat(Exception):
    def __init__(self, baris_ke, pesan):
        self.baris_ke = baris_ke
        super().__init__(f"Baris {baris_ke}: {pesan}")


def _baca_baris(sheet):
    header = [str(c.value).strip() if c.value else '' for c in sheet[1]]
    hilang = KOLOM_WAJIB - set(header)
    if hilang:
        raise ValueError(f"Kolom wajib hilang: {', '.join(sorted(hilang))}")

    for i, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        nilai = dict(zip(header, [c.value for c in row]))
        if all(v in (None, '') for v in nilai.values()):
            continue
        yield i, nilai


def _decimal(nilai, baris_ke, nama, wajib=True, default=Decimal('0')):
    if nilai in (None, ''):
        if wajib:
            raise BarisGalat(baris_ke, f"'{nama}' wajib diisi.")
        return default
    teks = str(nilai).strip()
    if ',' in teks and '.' in teks:
        teks = teks.replace('.', '').replace(',', '.')
    elif ',' in teks:
        teks = teks.replace(',', '.')
    try:
        return Decimal(teks)
    except InvalidOperation:
        raise BarisGalat(baris_ke, f"'{nama}' bukan angka valid: {nilai!r}")


def _ke_tanggal(nilai, baris_ke, nama, wajib=True):
    if nilai in (None, ''):
        if wajib:
            raise BarisGalat(baris_ke, f"'{nama}' wajib diisi.")
        return None
    if isinstance(nilai, datetime):
        return nilai.date()
    if isinstance(nilai, date):
        return nilai
    teks = str(nilai).strip()
    for fmt in FORMAT_TANGGAL:
        try:
            return datetime.strptime(teks, fmt).date()
        except ValueError:
            continue
    raise BarisGalat(baris_ke, f"'{nama}' bukan tanggal valid: {nilai!r}")


@transaction.atomic
def import_po_massal(file_obj, user=None):
    wb = load_workbook(file_obj, data_only=True)
    sheet = wb.active

    grup_rows = {}
    for baris_ke, nilai in _baca_baris(sheet):
        kode_grup = str(nilai.get('grup') or '').strip()
        if not kode_grup:
            raise BarisGalat(baris_ke, "'grup' wajib diisi untuk semua baris.")
        grup_rows.setdefault(kode_grup, []).append((baris_ke, nilai))

    berhasil, gagal = [], []

    for kode_grup, rows in grup_rows.items():
        sp = transaction.savepoint()
        try:
            po = _proses_satu_grup(kode_grup, rows, user=user)
            berhasil.append({"grup": kode_grup, "no_po": po.no_po, "id": po.id})
            transaction.savepoint_commit(sp)
        except (BarisGalat, ValueError, IntegrityError) as e:
            transaction.savepoint_rollback(sp)
            if isinstance(e, IntegrityError):
                pesan = (f"IntegrityError: {e}. Kemungkinan ada field wajib di "
                         f"DiauditModel yang belum di-set oleh import ini.")
            else:
                pesan = str(e)
            gagal.append({"grup": kode_grup, "pesan": pesan})

    return {"berhasil": berhasil, "gagal": gagal}


def _proses_satu_grup(kode_grup, rows, user=None):
    baris1_ke, header = rows[0]

    nama_entitas = str(header.get('entitas') or '').strip()
    nama_suplier = str(header.get('suplier') or '').strip()

    if not nama_entitas:
        raise BarisGalat(baris1_ke, "'entitas' wajib diisi pada baris pertama grup.")
    if not nama_suplier:
        raise BarisGalat(baris1_ke, "'suplier' wajib diisi pada baris pertama grup.")

    try:
        entitas = Entitas.objects.get(nama=nama_entitas)
    except Entitas.DoesNotExist:
        raise BarisGalat(baris1_ke, f"Entitas dengan nama '{nama_entitas}' tidak ditemukan.")
    except Entitas.MultipleObjectsReturned:
        raise BarisGalat(baris1_ke, f"Nama entitas '{nama_entitas}' tidak unik.")

    try:
        suplier = Suplier.objects.get(nama=nama_suplier)
    except Suplier.DoesNotExist:
        raise BarisGalat(baris1_ke, f"Suplier dengan nama '{nama_suplier}' tidak ditemukan.")
    except Suplier.MultipleObjectsReturned:
        raise BarisGalat(baris1_ke, f"Nama suplier '{nama_suplier}' tidak unik.")

    tanggal = _ke_tanggal(header.get('tanggal'), baris1_ke, 'tanggal')
    tanggal_kirim = _ke_tanggal(
        header.get('tanggal_kirim_diminta'), baris1_ke, 'tanggal_kirim_diminta', wajib=False
    )

    po = PurchaseOrder(
        entitas=entitas,
        suplier=suplier,
        tanggal=tanggal,
        tanggal_kirim_diminta=tanggal_kirim,
        ppn_persen=_decimal(header.get('ppn_persen'), baris1_ke, 'ppn_persen', wajib=False),
        catatan=str(header.get('catatan') or ''),
        status=StatusPO.DRAFT,
    )
    if user is not None and hasattr(po, 'dibuat_oleh'):
        po.dibuat_oleh = user
    po.save()

    for baris_ke, nilai in rows:
        nama_produk = str(nilai.get('produk') or '').strip()
        if not nama_produk:
            raise BarisGalat(baris_ke, "'produk' wajib diisi.")
        try:
            produk = Produk.objects.get(nama=nama_produk)
        except Produk.DoesNotExist:
            raise BarisGalat(baris_ke, f"Produk dengan nama '{nama_produk}' tidak ditemukan.")
        except Produk.MultipleObjectsReturned:
            raise BarisGalat(baris_ke, f"Nama produk '{nama_produk}' tidak unik.")

        qty_pesan = _decimal(nilai.get('qty_pesan'), baris_ke, 'qty_pesan')
        if qty_pesan <= 0:
            raise BarisGalat(baris_ke, "'qty_pesan' harus > 0.")

        harga = _decimal(nilai.get('harga_per_kg'), baris_ke, 'harga_per_kg')

        PurchaseOrderItem.objects.create(
            purchase_order=po,
            produk=produk,
            satuan=str(nilai.get('satuan') or 'kg'),
            qty_pesan=qty_pesan,
            harga_per_kg=harga,
        )

    return po